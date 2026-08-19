# script_nequi.py

import os
import re
import sys
import openpyxl
import pandas as pd
import pdfplumber

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


DELIMITADOR = "|||"


# ==============================================================================
# FUNCIONES AUXILIARES COMUNES
# ==============================================================================

def obtener_ruta_asset(ruta_relativa):
    """Obtiene la ruta absoluta para assets, funciona en desarrollo y PyInstaller."""
    if hasattr(sys, "_MEIPASS"):
        return os.path.join(sys._MEIPASS, ruta_relativa)

    return os.path.join(os.path.abspath("."), ruta_relativa)


def es_numero_financiero(texto):
    """
    Valida si un texto corresponde a un monto financiero.

    Ejemplos válidos:
        $-5,000.00
        $1,000.00
        $38.23
        -5000.00
        1000.00
    """

    if not texto:
        return False

    txt = (
        texto.strip()
        .replace("$", "")
        .replace(" ", "")
    )

    patron = r"^-?[\d\.,]+$"

    return bool(re.match(patron, txt))


def limpiar_monto(texto):
    """
    Limpia un monto para poder convertirlo posteriormente a número.
    """

    if not texto:
        return ""

    return (
        texto.strip()
        .replace("$", "")
        .replace(" ", "")
    )


def es_fecha_nequi(texto):
    """
    Detecta fechas con formato DD/MM/YYYY.

    Ejemplo:
        31/01/2026
    """

    if not texto:
        return False

    texto = texto.strip()

    return bool(
        re.match(
            r"^\d{2}/\d{2}/\d{4}$",
            texto
        )
    )


def limpiar_pie_pagina(words_fila):
    """
    Elimina textos correspondientes al pie de página de Nequi.
    """

    if not words_fila:
        return []

    linea = " ".join(
        w["text"]
        for w in words_fila
    ).upper()

    textos_pie = [
        "DEFENSOR CONSUMIDOR FINANCIERO",
        "VIGILADO",
        "SUPERINTENDENCIA FINANCIERA",
        "PRODUCTO PROTEGIDO",
        "FOGAFIN",
        "NEQUI GENERAN RENDIMIENTOS",
        "PUEDES CONSULTAR LA TASA",
    ]

    for texto in textos_pie:
        if texto in linea:
            return []

    return words_fila


def texto_delimitado_a_excel(
    lineas_texto,
    columnas,
    output_excel_path
):
    """
    Convierte líneas delimitadas en DataFrame y las exporta a Excel.
    """

    registros = []

    for linea in lineas_texto:

        partes = linea.split(DELIMITADOR)

        if len(partes) == len(columnas):

            registros.append(
                dict(
                    zip(
                        columnas,
                        partes
                    )
                )
            )

    df = pd.DataFrame(
        registros,
        columns=columnas
    )

    df.to_excel(
        output_excel_path,
        index=False
    )


# ==============================================================================
# PARSER DE EXTRACTOS NEQUI
# ==============================================================================

def extraer_lineas_extractos_nequi(
    pdf_path,
    progreso_callback=None
):
    """
    Extrae los movimientos de un extracto de Nequi.

    Estructura esperada:

        Fecha del movimiento | Descripción | Valor | Saldo

    Ejemplo:

        31/01/2026 | COMPRA PAQUETE PTM WOM | $-5,000.00 | $38.23
    """

    lineas_delimitadas = []

    patron_fecha = re.compile(
        r"^\d{2}/\d{2}/\d{4}$"
    )

    with pdfplumber.open(pdf_path) as pdf:

        total_paginas = len(pdf.pages)

        for indice_pagina, page in enumerate(
            pdf.pages,
            start=1
        ):

            ancho = page.width
            alto = page.height

            # ------------------------------------------------------------------
            # 1. EXTRAER PALABRAS
            # ------------------------------------------------------------------

            words = page.extract_words(
                x_tolerance=2,
                y_tolerance=3,
                keep_blank_chars=False
            )

            if not words:

                if progreso_callback:
                    progreso_callback(
                        indice_pagina,
                        total_paginas
                    )

                continue

            # ------------------------------------------------------------------
            # 2. DELIMITAR ÁREA DE MOVIMIENTOS
            #
            # El extracto mostrado tiene:
            #
            #   Resumen
            #   ...
            #   Fecha del movimiento
            #   Descripción
            #   Valor
            #   Saldo
            #   movimientos
            #   texto legal
            #
            # Por eso buscamos primero el encabezado de movimientos.
            # ------------------------------------------------------------------

            palabras_ordenadas = sorted(
                words,
                key=lambda w: (
                    w["top"],
                    w["x0"]
                )
            )

            y_inicio_movimientos = None
            y_fin_movimientos = alto

            for w in palabras_ordenadas:

                texto = w["text"].strip().upper()

                if (
                    "FECHA" in texto
                    and w["top"] > alto * 0.35
                ):
                    y_inicio_movimientos = w["top"]
                    break

            # Si no encontramos el encabezado,
            # utilizamos una posición aproximada.
            if y_inicio_movimientos is None:

                y_inicio_movimientos = alto * 0.50

            # ------------------------------------------------------------------
            # 3. BUSCAR EL FINAL DE LOS MOVIMIENTOS
            # ------------------------------------------------------------------

            textos_fin = [
                "DEFENSOR CONSUMIDOR",
                "NEQUI GENERAN",
                "PUEDES CONSULTAR",
                "VIGILADO",
                "SUPERINTENDENCIA",
            ]

            for w in palabras_ordenadas:

                if w["top"] <= y_inicio_movimientos:
                    continue

                texto = w["text"].strip().upper()

                for texto_fin in textos_fin:

                    if texto_fin in texto:

                        y_fin_movimientos = (
                            w["top"] - 3
                        )

                        break

                if y_fin_movimientos < alto:
                    break

            # ------------------------------------------------------------------
            # 4. OBTENER SOLO LAS PALABRAS DE LA TABLA
            # ------------------------------------------------------------------

            palabras_tabla = [
                w
                for w in words
                if (
                    w["top"] >= y_inicio_movimientos
                    and w["top"] < y_fin_movimientos
                )
            ]

            if not palabras_tabla:

                if progreso_callback:
                    progreso_callback(
                        indice_pagina,
                        total_paginas
                    )

                continue

            # ------------------------------------------------------------------
            # 5. DEFINIR COLUMNAS
            #
            # En el formato mostrado:
            #
            # FECHA       -> aproximadamente 10% - 28%
            # DESCRIPCIÓN -> aproximadamente 28% - 65%
            # VALOR       -> aproximadamente 65% - 82%
            # SALDO       -> aproximadamente 82% - 98%
            #
            # Se utilizan posiciones relativas para que funcione aunque
            # cambie ligeramente el tamaño de la página.
            # ------------------------------------------------------------------

            x_fecha_max = ancho * 0.30
            x_desc_max = ancho * 0.65
            x_valor_max = ancho * 0.82

            # ------------------------------------------------------------------
            # 6. ENCONTRAR FECHAS QUE FUNCIONAN COMO ANCLAS
            # ------------------------------------------------------------------

            fechas_anclaje = []

            for w in palabras_tabla:

                texto = w["text"].strip()

                if not patron_fecha.match(texto):
                    continue

                # La fecha debe estar ubicada en la zona izquierda.
                if w["x0"] >= x_fecha_max:
                    continue

                fechas_anclaje.append(
                    {
                        "fecha": texto,
                        "top": w["top"],
                        "word_obj": w
                    }
                )

            if not fechas_anclaje:

                if progreso_callback:
                    progreso_callback(
                        indice_pagina,
                        total_paginas
                    )

                continue

            # ------------------------------------------------------------------
            # 7. PROCESAR CADA MOVIMIENTO
            # ------------------------------------------------------------------

            for i in range(
                len(fechas_anclaje)
            ):

                fecha_actual = fechas_anclaje[i]

                y_inicio = (
                    fecha_actual["top"] - 2
                )

                if i + 1 < len(fechas_anclaje):

                    y_fin = (
                        fechas_anclaje[i + 1]["top"] - 2
                    )

                else:

                    y_fin = y_fin_movimientos

                palabras_bloque = [
                    w
                    for w in palabras_tabla
                    if (
                        w["top"] >= y_inicio
                        and w["top"] < y_fin
                    )
                ]

                movimiento = {
                    "FECHA": fecha_actual["fecha"],
                    "DESCRIPCIÓN": [],
                    "VALOR": "",
                    "SALDO": ""
                }

                # ------------------------------------------------------------------
                # 8. ORDENAR PALABRAS
                # ------------------------------------------------------------------

                palabras_bloque = sorted(
                    palabras_bloque,
                    key=lambda w: (
                        round(
                            w["top"] / 2.0
                        ) * 2.0,
                        w["x0"]
                    )
                )

                # ------------------------------------------------------------------
                # 9. CLASIFICAR CADA PALABRA
                # ------------------------------------------------------------------

                for w in palabras_bloque:

                    texto = w["text"].strip()

                    if not texto:
                        continue

                    if w is fecha_actual["word_obj"]:
                        continue

                    # Ignorar encabezados
                    texto_upper = texto.upper()

                    if texto_upper in [
                        "FECHA",
                        "DEL",
                        "MOVIMIENTO",
                        "DESCRIPCIÓN",
                        "VALOR",
                        "SALDO"
                    ]:
                        continue

                    x_mid = (
                        w["x0"] + w["x1"]
                    ) / 2.0

                    # ----------------------------------------------------------
                    # FECHA
                    # ----------------------------------------------------------

                    if x_mid < x_fecha_max:

                        # Si por alguna razón aparece otra fecha
                        # dentro del bloque, se ignora.
                        if patron_fecha.match(texto):
                            continue

                        movimiento[
                            "DESCRIPCIÓN"
                        ].append(texto)

                    # ----------------------------------------------------------
                    # DESCRIPCIÓN
                    # ----------------------------------------------------------

                    elif x_mid < x_desc_max:

                        movimiento[
                            "DESCRIPCIÓN"
                        ].append(texto)

                    # ----------------------------------------------------------
                    # VALOR
                    # ----------------------------------------------------------

                    elif x_mid < x_valor_max:

                        if (
                            es_numero_financiero(texto)
                            and not movimiento["VALOR"]
                        ):

                            movimiento[
                                "VALOR"
                            ] = texto

                    # ----------------------------------------------------------
                    # SALDO
                    # ----------------------------------------------------------

                    else:

                        if (
                            es_numero_financiero(texto)
                            and not movimiento["SALDO"]
                        ):

                            movimiento[
                                "SALDO"
                            ] = texto

                # ------------------------------------------------------------------
                # 10. CONSTRUIR DESCRIPCIÓN
                # ------------------------------------------------------------------

                descripcion = " ".join(
                    movimiento["DESCRIPCIÓN"]
                ).strip()

                # ------------------------------------------------------------------
                # 11. VALIDAR QUE REALMENTE SEA UN MOVIMIENTO
                # ------------------------------------------------------------------

                if not descripcion:
                    continue

                if not movimiento["VALOR"]:
                    continue

                if not movimiento["SALDO"]:
                    continue

                # Evitar encabezados
                if (
                    "DESCRIPCI" in descripcion.upper()
                    or "FECHA DEL MOVIMIENTO"
                    in descripcion.upper()
                ):
                    continue

                registro = [
                    movimiento["FECHA"],
                    descripcion,
                    movimiento["VALOR"],
                    movimiento["SALDO"]
                ]

                lineas_delimitadas.append(
                    DELIMITADOR.join(registro)
                )

            # ------------------------------------------------------------------
            # 12. PROGRESO
            # ------------------------------------------------------------------

            if progreso_callback:

                progreso_callback(
                    indice_pagina,
                    total_paginas
                )

    return lineas_delimitadas


# ==============================================================================
# REORGANIZACIÓN DEL EXCEL
# ==============================================================================

def reorganizar_excel(excel_path):
    """
    Reorganiza el Excel generado por el extractor de Nequi.
    """

    if not os.path.exists(excel_path):
        raise FileNotFoundError(
            f"No se encontró el archivo: {excel_path}"
        )

    df = pd.read_excel(
        excel_path,
        dtype=str
    )

    if df.empty:
        return excel_path

    if "VALOR" not in df.columns:
        return excel_path

    # --------------------------------------------------------------------------
    # LIMPIAR VALORES
    # --------------------------------------------------------------------------

    valores_numericos = pd.to_numeric(
        df["VALOR"]
        .astype(str)
        .str.replace(",", "", regex=False)
        .str.replace("$", "", regex=False)
        .str.replace(" ", "", regex=False)
        .str.strip(),
        errors="coerce"
    ).fillna(0.0)

    saldos_numericos = pd.to_numeric(
        df["SALDO"]
        .astype(str)
        .str.replace(",", "", regex=False)
        .str.replace("$", "", regex=False)
        .str.replace(" ", "", regex=False)
        .str.strip(),
        errors="coerce"
    ).fillna(0.0)

    # --------------------------------------------------------------------------
    # TOTALES
    # --------------------------------------------------------------------------

    cargos_sum = valores_numericos[
        valores_numericos < 0
    ].sum()

    abonos_sum = valores_numericos[
        valores_numericos > 0
    ].sum()

    saldo_anterior = (
        saldos_numericos.iloc[0]
        - valores_numericos.iloc[0]
    )

    saldo_actual = saldos_numericos.iloc[-1]

    df_resumen = pd.DataFrame(
        [
            {
                "CONCEPTO": "SALDO ANTERIOR",
                "MONTO": saldo_anterior
            },
            {
                "CONCEPTO": "TOTAL CARGOS",
                "MONTO": cargos_sum
            },
            {
                "CONCEPTO": "TOTAL ABONOS",
                "MONTO": abonos_sum
            },
            {
                "CONCEPTO": "SALDO ACTUAL",
                "MONTO": saldo_actual
            }
        ]
    )

    # --------------------------------------------------------------------------
    # MES
    # --------------------------------------------------------------------------

    fechas_dt = pd.to_datetime(
        df["FECHA"],
        errors="coerce",
        dayfirst=True,
        format="mixed"
    )

    df["MES"] = (
        fechas_dt.dt.month
        .fillna(0)
        .astype(int)
    )

    df["_VALOR_NUM"] = valores_numericos

    df["_CARGOS_TEMP"] = df[
        "_VALOR_NUM"
    ].apply(
        lambda x: x if x < 0 else 0.0
    )

    df["_ABONOS_TEMP"] = df[
        "_VALOR_NUM"
    ].apply(
        lambda x: x if x > 0 else 0.0
    )

    # --------------------------------------------------------------------------
    # AGRUPAR POR MES Y DESCRIPCIÓN
    # --------------------------------------------------------------------------

    df_grouped = (
        df.groupby(
            ["MES", "DESCRIPCIÓN"],
            as_index=False
        )
        .agg(
            CARGOS=(
                "_CARGOS_TEMP",
                "sum"
            ),
            ABONOS=(
                "_ABONOS_TEMP",
                "sum"
            ),
            NETO=(
                "_VALOR_NUM",
                "sum"
            )
        )
        .sort_values(
            by=[
                "MES",
                "CARGOS"
            ],
            ascending=[
                True,
                True
            ]
        )
    )

    # --------------------------------------------------------------------------
    # TOTAL GENERAL
    # --------------------------------------------------------------------------

    fila_total = {
        "MES": "",
        "DESCRIPCIÓN": "TOTAL GENERAL",
        "CARGOS": float(
            df_grouped["CARGOS"].sum()
        ),
        "ABONOS": float(
            df_grouped["ABONOS"].sum()
        ),
        "NETO": float(
            df_grouped["NETO"].sum()
        )
    }

    df_conceptos = pd.concat(
        [
            df_grouped,
            pd.DataFrame([fila_total])
        ],
        ignore_index=True
    )

    # --------------------------------------------------------------------------
    # ORDENAR DATOS
    # NEGATIVOS PRIMERO
    # --------------------------------------------------------------------------

    df_datos = df.copy()

    df_datos["VALOR"] = valores_numericos
    df_datos["SALDO"] = saldos_numericos

    df_datos["_ORDEN"] = (
        df_datos["VALOR"]
        >= 0
    ).astype(int)

    df_datos = (
        df_datos
        .sort_values(
            by="_ORDEN",
            kind="stable"
        )
        .drop(
            columns=[
                "_ORDEN",
                "MES",
                "_VALOR_NUM",
                "_CARGOS_TEMP",
                "_ABONOS_TEMP"
            ],
            errors="ignore"
        )
    )

    # --------------------------------------------------------------------------
    # EXPORTAR
    # --------------------------------------------------------------------------

    with pd.ExcelWriter(
        excel_path,
        engine="openpyxl"
    ) as writer:

        df_datos.to_excel(
            writer,
            sheet_name="Datos",
            index=False
        )

        df_resumen.to_excel(
            writer,
            sheet_name="Resumen",
            index=False
        )

        df_conceptos.to_excel(
            writer,
            sheet_name="Conceptos",
            index=False
        )

        workbook = writer.book

        FORMATO_MONEDA = (
            '"$"#,##0.00;[Red]("$"#,##0.00);"-"'
        )

        for sheet_name in workbook.sheetnames:

            worksheet = workbook[
                sheet_name
            ]

            for col in worksheet.iter_cols(
                1,
                worksheet.max_column
            ):

                header = (
                    str(
                        col[0].value
                    ).upper()
                    if col[0].value
                    else ""
                )

                if header in [
                    "VALOR",
                    "SALDO",
                    "CARGOS",
                    "ABONOS",
                    "NETO",
                    "MONTO"
                ]:

                    for cell in col[1:]:

                        if (
                            cell.value is not None
                            and isinstance(
                                cell.value,
                                (
                                    int,
                                    float
                                )
                            )
                        ):

                            cell.number_format = (
                                FORMATO_MONEDA
                            )

    return excel_path


# ==============================================================================
# ORQUESTADOR PRINCIPAL
# ==============================================================================

def ejecutar_proceso_exportacion(
    pdf_path,
    output_excel_path=None,
    progreso_callback=None
):
    """
    Procesa un extracto de Nequi y genera:

        1. Archivo TXT plano
        2. Archivo XLSX
        3. Hoja Datos
        4. Hoja Resumen
        5. Hoja Conceptos
    """

    columnas = [
        "FECHA",
        "DESCRIPCIÓN",
        "VALOR",
        "SALDO"
    ]

    # --------------------------------------------------------------------------
    # 1. EXTRAER MOVIMIENTOS
    # --------------------------------------------------------------------------

    lineas_plana = (
        extraer_lineas_extractos_nequi(
            pdf_path,
            progreso_callback
        )
    )

    if not lineas_plana:

        raise ValueError(
            "No se pudieron extraer movimientos del PDF de Nequi."
        )

    # --------------------------------------------------------------------------
    # 2. DEFINIR NOMBRE DE SALIDA
    # --------------------------------------------------------------------------

    if not output_excel_path:

        base_path, _ = os.path.splitext(
            pdf_path
        )

        output_excel_path = (
            f"{base_path}_convertido.xlsx"
        )

    # --------------------------------------------------------------------------
    # 3. ARCHIVO TXT
    # --------------------------------------------------------------------------

    base_txt_path, _ = os.path.splitext(
        output_excel_path
    )

    txt_path = (
        f"{base_txt_path}_plano.txt"
    )

    encabezado_txt = DELIMITADOR.join(
        columnas
    )

    contenido_txt = [
        encabezado_txt
    ] + lineas_plana

    with open(
        txt_path,
        "w",
        encoding="utf-8"
    ) as f:

        f.write(
            "\n".join(
                contenido_txt
            )
        )

    # --------------------------------------------------------------------------
    # 4. TXT -> EXCEL
    # --------------------------------------------------------------------------

    texto_delimitado_a_excel(
        lineas_plana,
        columnas,
        output_excel_path
    )

    # --------------------------------------------------------------------------
    # 5. REORGANIZAR EXCEL
    # --------------------------------------------------------------------------

    reorganizar_excel(
        output_excel_path
    )

    return output_excel_path