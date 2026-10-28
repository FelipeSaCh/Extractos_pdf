import os
import re
import pdfplumber
import pandas as pd
from openpyxl.styles import Font
import unicodedata

DELIMITADOR = "|||"
FORMATO_MONEDA = '"$"#,##0.00;[Red]("$"#,##0.00);"-"'
COLUMNAS_MONEDA = ("VALOR", "SALDO", "CARGOS", "ABONOS", "NETO", "MONTO")
MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO",
         "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]

_ABREV_MES = {m[:3].lower(): i for i, m in enumerate(MESES, start=1)}
_PATRON_FINANCIERO = re.compile(r"^-?[\d.,]+$")
_PATRON_PAG = re.compile(r"(?i)^p[áa]g(?:ina)?\.?$")
_SECUENCIA_PAG = (re.compile(r"^\d+$"), re.compile(r"(?i)^de\.?$"), re.compile(r"^\d+$"))

HOJA_INFO = "Info"
CLAVE_CUENTA = "CUENTA BANCARIA"

FRASES_CUENTA = (
    "NUMERO DE CUENTA", "NUMERO DE LA CUENTA", "NO. DE CUENTA", "NO DE CUENTA", "NRO. CUENTA", "NRO CUENTA",
    "CUENTA NO.", "CUENTA NO", "CUENTA NRO.", "CUENTA NRO", "CUENTA N°", "CUENTA #",
    "CUENTA DE AHORROS", "CUENTA AHORROS", "CUENTA CORRIENTE", "CTA. AHORROS", "CTA AHORROS",
    "CTA. CORRIENTE", "CTA CORRIENTE", "CTA.", "CTA", "NUMERO DE PRODUCTO", "PRODUCTO","NÚMERO",
)
_NUMERO_CUENTA = r"[^\d\n]{0,25}\n?[^\d\n]{0,10}(\d[\d -]{4,24}\d)"


def _sin_tildes(texto):
    return unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode().upper()


def _patron_frase(frase):
    palabras = [re.escape(p) for p in _sin_tildes(frase).split()]
    return re.compile(r"(?<![A-Z])" + r"\s*".join(palabras) + r"(?![A-Z])" + _NUMERO_CUENTA)


_PATRONES_CUENTA = [_patron_frase(f) for f in FRASES_CUENTA]


def buscar_cuenta_bancaria(texto, frases=None):
    texto = _sin_tildes(texto or "")
    for patron in [_patron_frase(f) for f in frases or ()] + _PATRONES_CUENTA:
        for m in patron.finditer(texto):
            numero = re.sub(r"\s", "", m[1])
            if sum(c.isdigit() for c in numero) >= 6:
                return numero
    return ""


def extraer_cuenta_bancaria(pdf_path, frases=None):
    try:
        with pdfplumber.open(pdf_path) as pdf:
            texto = "\n".join(p.extract_text() or "" for p in pdf.pages[:2])
    except Exception:
        return ""
    return buscar_cuenta_bancaria(texto, frases)

def obtener_ruta_asset(ruta_relativa):
    return os.path.join(os.path.dirname(os.path.abspath(__file__)), ruta_relativa)


def es_numero_financiero(texto):
    return bool(texto) and bool(_PATRON_FINANCIERO.match(texto.strip().replace("$", "").replace(" ", "")))


def limpiar_pie_pagina(words_fila, frases=()):
    if frases:
        linea = " ".join(w["text"] for w in words_fila).upper()
        if any(f in linea for f in frases):
            return []
    textos = [w["text"].strip() for w in words_fila]
    quitar, i = set(), 0
    while i < len(textos):
        if not _PATRON_PAG.match(textos[i]):
            i += 1
            continue
        j = i + 1
        for patron in _SECUENCIA_PAG:
            if j >= len(textos) or not patron.match(textos[j]):
                break
            j += 1
        quitar.update(range(i, j))
        i = j
    return [w for k, w in enumerate(words_fila) if k not in quitar] if quitar else words_fila


def agrupar_filas(words, paso):
    filas = {}
    for w in words:
        filas.setdefault(round(w["top"] / paso) * paso, []).append(w)
    return [sorted(filas[y], key=lambda w: w["x0"]) for y in sorted(filas)]


def texto_delimitado_a_excel(lineas_texto, columnas, output_excel_path):
    partes = (linea.split(DELIMITADOR) for linea in lineas_texto)
    registros = [dict(zip(columnas, p)) for p in partes if len(p) == len(columnas)]
    pd.DataFrame(registros, columns=columnas).to_excel(output_excel_path, index=False)


def a_numero(serie):
    return pd.to_numeric(serie.astype(str).str.replace(r"[,$\s]", "", regex=True), errors="coerce").fillna(0.0)


def mes_de_fechas(fechas):
    texto = fechas.fillna("").astype(str).str.strip()

    def numero(patron):
        return pd.to_numeric(texto.str.extract(patron, expand=False), errors="coerce")

    abreviado = texto.str.lower().str.extract(rf"\b({'|'.join(_ABREV_MES)})", expand=False).map(_ABREV_MES)
    mes = numero(r"^\d{4}[/.-](\d{1,2})").fillna(numero(r"^\d{1,2}[/.-](\d{1,2})")).fillna(abreviado)
    return mes.where(mes.between(1, 12)).fillna(0).astype(int)


def _resumen(valores, saldos=None, saldo_anterior=None, saldo_actual=None):
    conciliar = saldo_anterior is not None and saldo_actual is not None
    if saldos is not None and not saldos.empty:
        saldo_anterior = saldos.iloc[0] - valores.iloc[0] if saldo_anterior is None else saldo_anterior
        saldo_actual = saldos.iloc[-1] if saldo_actual is None else saldo_actual
    cargos, abonos = valores[valores < 0].sum(), valores[valores > 0].sum()
    filas = [("SALDO ANTERIOR", saldo_anterior)] if saldo_anterior is not None else []
    filas += [("TOTAL CARGOS", cargos), ("TOTAL ABONOS", abonos)]
    if saldo_actual is not None:
        filas.append(("SALDO ACTUAL", saldo_actual))
    if conciliar:
        filas.append(("DIFERENCIA CONCILIACIÓN", round(saldo_anterior + cargos + abonos - saldo_actual, 2)))
    return pd.DataFrame(filas, columns=["CONCEPTO", "MONTO"])


def _formato_moneda(ws):
    for col in ws.iter_cols(1, ws.max_column):
        if str(col[0].value or "").upper() in COLUMNAS_MONEDA:
            for celda in col[1:]:
                if isinstance(celda.value, (int, float)):
                    celda.number_format = FORMATO_MONEDA


def _anexar_resumenes(ws, bloques, col_concepto, col_monto):
    negrita = Font(bold=True)
    fila = ws.max_row + 3
    for titulo, resumen in bloques:
        ws.cell(fila, col_concepto, titulo).font = negrita
        for i, (concepto, monto) in enumerate(resumen.itertuples(index=False), start=1):
            ws.cell(fila + i, col_concepto, concepto).font = negrita
            celda = ws.cell(fila + i, col_monto, float(monto))
            celda.number_format = FORMATO_MONEDA
            celda.font = negrita
        fila += len(resumen) + 3


def _bloque_resumen(columnas, col_desc, titulo, resumen):
    return pd.DataFrame({
        col_desc: [titulo, *resumen["CONCEPTO"]],
        "VALOR": [None, *resumen["MONTO"]],
    }).reindex(columns=columnas)


def _intercalar_resumenes(columnas, col_desc, bloques, resumen_general):
    partes, negritas, fila = [], [], 2

    def agregar(df, negrita=False):
        nonlocal fila
        partes.append(df)
        if negrita:
            negritas.extend(range(fila, fila + len(df)))
        fila += len(df)

    for titulo, resumen, filas in bloques:
        agregar(filas)
        agregar(pd.DataFrame(index=range(1), columns=columnas))
        agregar(_bloque_resumen(columnas, col_desc, titulo, resumen), True)
        agregar(pd.DataFrame(index=range(2), columns=columnas))
    agregar(_bloque_resumen(columnas, col_desc, "RESUMEN GENERAL", resumen_general), True)
    return pd.concat(partes, ignore_index=True), negritas


def reorganizar_excel(excel_path, saldo_anterior=None, saldo_actual=None, cuenta_bancaria=None):
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"No se encontró el archivo: {excel_path}")
    with pd.ExcelFile(excel_path) as xl:
        df = pd.read_excel(xl, sheet_name="Datos" if "Datos" in xl.sheet_names else 0, dtype=str)
        if cuenta_bancaria is None and HOJA_INFO in xl.sheet_names:
            info = pd.read_excel(xl, sheet_name=HOJA_INFO, dtype=str).fillna("")
            cuenta_bancaria = next(iter(info.loc[info["CLAVE"] == CLAVE_CUENTA, "VALOR"]), "")

    col_desc = "DESCRIPCIÓN" if "DESCRIPCIÓN" in df.columns else "DESCRIPCION"
    valores = a_numero(df["VALOR"])
    tiene_saldo = "SALDO" in df.columns and df["SALDO"].fillna("").str.strip().ne("").any()
    saldos = a_numero(df["SALDO"]) if tiene_saldo else None
    meses = mes_de_fechas(df["FECHA"]) if "FECHA" in df.columns else pd.Series(0, index=df.index)
    resumen = _resumen(valores, saldos, saldo_anterior, saldo_actual)

    calc = pd.DataFrame({
        "MES": meses,
        col_desc: df[col_desc].fillna(""),
        "CARGOS": valores.where(valores < 0, 0.0),
        "ABONOS": valores.where(valores > 0, 0.0),
        "NETO": valores,
    })
    agrupado = calc.groupby(["MES", col_desc], as_index=False).sum().sort_values(["MES", "CARGOS"])
    vacia = dict.fromkeys(["MES", col_desc, "CARGOS", "ABONOS", "NETO"])
    filas = []
    for i, (_, grupo) in enumerate(agrupado.groupby("MES", sort=True)):
        filas += [dict(vacia) for _ in range(3)] if i else []
        filas += grupo.to_dict("records")
    filas += [dict(vacia) for _ in range(3)]
    filas.append({**vacia, col_desc: "TOTAL GENERAL", "CARGOS": float(agrupado["CARGOS"].sum()),
                  "ABONOS": float(agrupado["ABONOS"].sum()), "NETO": float(agrupado["NETO"].sum())})
    conceptos = pd.DataFrame(filas, columns=list(vacia))

    datos = df.assign(VALOR=valores)
    if tiene_saldo:
        datos["SALDO"] = saldos
    orden_meses = list(pd.unique(meses))
    negritas = []
    if len(orden_meses) > 1:
        bloques = [
            (f"RESUMEN {MESES[mes - 1]}" if mes else "RESUMEN SIN FECHA",
             _resumen(valores[meses == mes], saldos[meses == mes] if tiene_saldo else None),
             datos[meses == mes])
            for mes in orden_meses
        ]
        datos, negritas = _intercalar_resumenes(datos.columns, col_desc, bloques, resumen)
    else:
        datos = datos.loc[(valores >= 0).sort_values(kind="stable").index]

    with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
        datos.to_excel(writer, sheet_name="Datos", index=False)
        resumen.to_excel(writer, sheet_name="Resumen", index=False)
        conceptos.to_excel(writer, sheet_name="Conceptos", index=False)
        if cuenta_bancaria:
            pd.DataFrame({"CLAVE": [CLAVE_CUENTA], "VALOR": [str(cuenta_bancaria)]}).to_excel(
                writer, sheet_name=HOJA_INFO, index=False)
        libro = writer.book
        for ws in libro.worksheets:
            _formato_moneda(ws)
        col_concepto, col_monto = datos.columns.get_loc(col_desc) + 1, datos.columns.get_loc("VALOR") + 1
        if negritas:
            negrita = Font(bold=True)
            for fila in negritas:
                libro["Datos"].cell(fila, col_concepto).font = negrita
                libro["Datos"].cell(fila, col_monto).font = negrita
        else:
            _anexar_resumenes(libro["Datos"], [("RESUMEN", resumen)], col_concepto, col_monto)
        _anexar_resumenes(libro["Conceptos"], [("RESUMEN", resumen)],
                          conceptos.columns.get_loc(col_desc) + 1, conceptos.columns.get_loc("NETO") + 1)
    return excel_path


def exportar(lineas, columnas, ruta_origen, output_excel_path=None, saldo_anterior=None, saldo_actual=None,
             frases_cuenta=None):
    output_excel_path = output_excel_path or f"{os.path.splitext(ruta_origen)[0]}_convertido.xlsx"
    with open(f"{os.path.splitext(output_excel_path)[0]}_plano.txt", "w", encoding="utf-8") as f:
        f.write("\n".join([DELIMITADOR.join(columnas), *lineas]))
    texto_delimitado_a_excel(lineas, columnas, output_excel_path)
    reorganizar_excel(output_excel_path, saldo_anterior, saldo_actual,
                      extraer_cuenta_bancaria(ruta_origen, frases_cuenta))
    return output_excel_path