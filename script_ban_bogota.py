import os
import re
import sys
import openpyxl
import pandas as pd
import pdfplumber
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DELIMITADOR = "|||"


# ==============================================================================
# FUNCIONES AUXILIARES COMUNES
# ==============================================================================
def obtener_ruta_asset(ruta_relativa):
    """Obtiene la ruta absoluta para assets, funciona en desarrollo y en el .exe de PyInstaller."""
    if hasattr(sys, "_MEIPASS"):
        return os.path.join(sys._MEIPASS, ruta_relativa)
    return os.path.join(os.path.abspath("."), ruta_relativa)


def es_numero_financiero(texto):
    """Valida si un texto es un monto monetario (positivo, negativo o decimales)."""
    if not texto:
        return False
    txt = texto.strip().replace("$", "").replace(" ", "")
    patron = r"^-?[\d\.,]+$"
    return bool(re.match(patron, txt))


_PATRON_PAG_TOKEN = re.compile(r"(?i)^p[áa]g(?:ina)?\.?$")
_PATRON_DE_TOKEN = re.compile(r"(?i)^de\.?$")
_PATRON_NUM_TOKEN = re.compile(r"^\d+$")


def limpiar_pie_pagina(words_fila):
    """Elimina 'Página N de M' y avisos legales de pie de página."""
    linea_texto = " ".join([w["text"] for w in words_fila]).upper()
    if (
        "NOTIFICAR CUALQUIER REPARO" in linea_texto
        or "REVISORÍA FISCAL" in linea_texto
        or "VALOR PENDIENTE POR APLICAR" in linea_texto
    ):
        return []

    n = len(words_fila)
    idx_a_quitar = set()
    i = 0
    while i < n:
        if _PATRON_PAG_TOKEN.match(words_fila[i]["text"].strip()):
            seq = [i]
            j = i + 1
            if j < n and _PATRON_NUM_TOKEN.match(words_fila[j]["text"].strip()):
                seq.append(j)
                j += 1
                if j < n and _PATRON_DE_TOKEN.match(
                    words_fila[j]["text"].strip()
                ):
                    seq.append(j)
                    j += 1
                    if j < n and _PATRON_NUM_TOKEN.match(
                        words_fila[j]["text"].strip()
                    ):
                        seq.append(j)
                        j += 1
            idx_a_quitar.update(seq)
            i = j
        else:
            i += 1

    if not idx_a_quitar:
        return words_fila
    return [w for k, w in enumerate(words_fila) if k not in idx_a_quitar]


def texto_delimitado_a_excel(lineas_texto, columnas, output_excel_path):
    """Convierte las líneas delimitadas a un DataFrame y lo exporta a Excel."""
    registros = []
    for linea in lineas_texto:
        partes = linea.split(DELIMITADOR)
        if len(partes) == len(columnas):
            registros.append(dict(zip(columnas, partes)))

    df = pd.DataFrame(registros, columns=columnas)
    df.to_excel(output_excel_path, index=False)


# ==============================================================================
# PARSER EXTRACCIÓN MILIMÉTRICA BANCO DE BOGOTÁ
# ==============================================================================


def extraer_lineas_extractos_pyme(pdf_path):
    patron_fecha = r"^\b\d{1,2}/\d{2}\b"
    lineas_delimitadas = []

    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            anc = page.width
            alt = page.height

            # 1. DELIMITACIÓN ESTRICTA DEL ÁREA ÚTIL (Crop Zone)
            # Ignora el margen izquierdo (texto vertical) y los márgenes exterior/pie
            x_min_tabla = anc * 0.05  # Deja por fuera "VIGILADO..."
            x_max_tabla = anc * 0.98

            # Filtrar palabras dentro de los márgenes horizontales de la tabla
            words = [
                w
                for w in page.extract_words()
                if x_min_tabla <= w["x0"] <= x_max_tabla
            ]
            if not words:
                continue

            # Columnas basadas en X relativas dentro del área útil
            x_fecha_max = anc * 0.11
            x_cod_max = anc * 0.16
            x_desc_max = anc * 0.43
            x_ciudad_max = anc * 0.53
            x_oficina_max = anc * 0.67
            x_doc_max = anc * 0.74
            x_valor_max = anc * 0.88

            # 2. ENCONTRAR LÍMITE INFERIOR (Corte antes del pie de página)
            y_max_tabla = alt
            for w in sorted(words, key=lambda x: x["top"]):
                txt_upper = w["text"].upper()
                if (
                    "FIN MOVIMIENTOS" in txt_upper
                    or "NOTIFICAR" in txt_upper
                    or "REVISORÍA" in txt_upper
                ):
                    y_max_tabla = w["top"] - 2
                    break

            # 3. ANCLAJE POR FECHAS (Columna 1 real)
            fechas_anclaje = []
            for w in sorted(words, key=lambda x: (x["top"], x["x0"])):
                if w["top"] >= y_max_tabla:
                    continue

                txt = w["text"].strip()
                # La fecha debe estar en la primera columna
                if w["x0"] < x_fecha_max and re.match(patron_fecha, txt):
                    fechas_anclaje.append(
                        {
                            "fecha": txt,
                            "top": w["top"],
                            "word_obj": w,
                        }
                    )

            if not fechas_anclaje:
                continue

            # 4. PROCESAR CADA BLOQUE DE FECHA
            for i in range(len(fechas_anclaje)):
                f_actual = fechas_anclaje[i]
                y_inicio = f_actual["top"] - 1.5

                if i + 1 < len(fechas_anclaje):
                    y_fin = fechas_anclaje[i + 1]["top"] - 1.5
                else:
                    y_fin = y_max_tabla

                # Palabras dentro del bloque vertical Y y de la tabla X
                palabras_bloque = [
                    w
                    for w in words
                    if y_inicio <= w["top"] < y_fin and w["top"] < y_max_tabla
                ]

                tx_actual = {
                    "FECHA": f_actual["fecha"],
                    "COD_TRANS": [],
                    "DESCRIPCIÓN": [],
                    "CIUDAD": [],
                    "OFICINA": [],
                    "DOCUMENTO": [],
                    "VALOR": "",
                    "SALDO": "",
                }

                palabras_ordenadas = sorted(
                    palabras_bloque,
                    key=lambda w: (round(w["top"] / 2.0) * 2.0, w["x0"]),
                )

                for w in palabras_ordenadas:
                    txt = w["text"].strip()
                    if not txt or set(txt) == {"-"}:
                        continue

                    if w == f_actual["word_obj"]:
                        continue

                    x_mid = (w["x0"] + w["x1"]) / 2.0

                    if x_mid < x_cod_max:
                        tx_actual["COD_TRANS"].append(txt)
                    elif x_mid < x_desc_max:
                        tx_actual["DESCRIPCIÓN"].append(txt)
                    elif x_mid < x_ciudad_max:
                        tx_actual["CIUDAD"].append(txt)
                    elif x_mid < x_oficina_max:
                        tx_actual["OFICINA"].append(txt)
                    elif x_mid < x_doc_max:
                        tx_actual["DOCUMENTO"].append(txt)
                    elif x_mid < x_valor_max:
                        if (
                            es_numero_financiero(txt)
                            and not tx_actual["VALOR"]
                        ):
                            tx_actual["VALOR"] = txt
                    else:
                        if (
                            es_numero_financiero(txt)
                            and not tx_actual["SALDO"]
                        ):
                            tx_actual["SALDO"] = txt

                desc = " ".join(tx_actual["DESCRIPCIÓN"]).strip()

                if (
                    "DESCRIPCI" in desc.upper()
                    or "FECHA" in tx_actual["FECHA"].upper()
                ):
                    continue

                registro = [
                    tx_actual["FECHA"],
                    " ".join(tx_actual["COD_TRANS"]).strip(),
                    desc,
                    " ".join(tx_actual["CIUDAD"]).strip(),
                    " ".join(tx_actual["OFICINA"]).strip(),
                    " ".join(tx_actual["DOCUMENTO"]).strip(),
                    tx_actual["VALOR"],
                    tx_actual["SALDO"],
                ]
                lineas_delimitadas.append(DELIMITADOR.join(registro))

    return lineas_delimitadas

# ==============================================================================
# REORGANIZACIÓN Y GENERACIÓN DE REPORTES EN EXCEL
# ==============================================================================


def reorganizar_excel(excel_path):
    """Calcula totales, conciliación y conceptos."""
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"No se encontró el archivo: {excel_path}")

    xl = pd.ExcelFile(excel_path)
    sheet_a_leer = "Datos" if "Datos" in xl.sheet_names else 0
    df = pd.read_excel(excel_path, sheet_name=sheet_a_leer, dtype=str)

    if "VALOR" not in df.columns or df.empty:
        return excel_path

    col_desc = "DESCRIPCIÓN" if "DESCRIPCIÓN" in df.columns else "DESCRIPCION"

    # PASO 1: LIMPIEZA DE NÚMEROS
    valores_numericos = pd.to_numeric(
        df["VALOR"]
        .astype(str)
        .str.replace(",", "", regex=False)
        .str.replace("$", "", regex=False)
        .str.replace(" ", "", regex=False)
        .str.strip(),
        errors="coerce",
    ).fillna(0.0)

    tiene_columna_saldo = "SALDO" in df.columns
    saldos_numericos = pd.Series(dtype=float)

    if tiene_columna_saldo:
        saldos_numericos = pd.to_numeric(
            df["SALDO"]
            .astype(str)
            .str.replace(",", "", regex=False)
            .str.replace("$", "", regex=False)
            .str.replace(" ", "", regex=False)
            .str.strip(),
            errors="coerce",
        ).fillna(0.0)

    # PASO 2: CÁLCULOS SOBRE EL DF ORIGINAL Y EXTRACCIÓN DE MES
    cargos_sum = valores_numericos[valores_numericos < 0].sum()
    abonos_sum = valores_numericos[valores_numericos > 0].sum()

    filas_resumen = []
    if tiene_columna_saldo and not saldos_numericos.empty:
        saldo_anterior = saldos_numericos.iloc[0] - valores_numericos.iloc[0]
        saldo_actual = saldos_numericos.iloc[-1]
        filas_resumen.append(
            {"CONCEPTO": "SALDO ANTERIOR", "MONTO": saldo_anterior}
        )

    filas_resumen.append({"CONCEPTO": "TOTAL CARGOS", "MONTO": cargos_sum})
    filas_resumen.append({"CONCEPTO": "TOTAL ABONOS", "MONTO": abonos_sum})

    if tiene_columna_saldo and not saldos_numericos.empty:
        filas_resumen.append(
            {"CONCEPTO": "SALDO ACTUAL", "MONTO": saldo_actual}
        )

    df_conciliacion = pd.DataFrame(filas_resumen)

    # EXTRAER MES
    df_calc = df.copy()
    fechas_dt = pd.to_datetime(
        df_calc["FECHA"], errors="coerce", dayfirst=True, format="mixed"
    )
    meses_extraidos = fechas_dt.dt.month.fillna(
        df_calc["FECHA"].astype(str).str.extract(r"[/.-](\d{1,2})", expand=False)
    )

    df_calc["MES"] = (
        pd.to_numeric(meses_extraidos, errors="coerce").fillna(0).astype(int)
    )
    df_calc["_VALOR_NUM"] = valores_numericos
    df_calc["_CARGOS_TEMP"] = df_calc["_VALOR_NUM"].apply(
        lambda x: x if x < 0 else 0.0
    )
    df_calc["_ABONOS_TEMP"] = df_calc["_VALOR_NUM"].apply(
        lambda x: x if x > 0 else 0.0
    )

    # Agrupar por MES y DESCRIPCIÓN
    df_grouped = (
        df_calc.groupby(["MES", col_desc], as_index=False)
        .agg(
            CARGOS=("_CARGOS_TEMP", "sum"),
            ABONOS=("_ABONOS_TEMP", "sum"),
            NETO=("_VALOR_NUM", "sum"),
        )
        .sort_values(by=["MES", "CARGOS"], ascending=[True, True])
    )

    filas_con_espacios = []
    meses_unicos = df_grouped["MES"].unique()

    fila_vacia = {
        "MES": "",
        col_desc: "",
        "CARGOS": None,
        "ABONOS": None,
        "NETO": None,
    }

    for i, mes in enumerate(meses_unicos):
        df_mes = df_grouped[df_grouped["MES"] == mes]
        filas_con_espacios.extend(df_mes.to_dict("records"))

        if i < len(meses_unicos) - 1:
            filas_con_espacios.extend([fila_vacia.copy() for _ in range(3)])

    filas_con_espacios.extend([fila_vacia.copy() for _ in range(3)])

    fila_total = {
        "MES": "",
        col_desc: "TOTAL GENERAL",
        "CARGOS": float(df_grouped["CARGOS"].sum()),
        "ABONOS": float(df_grouped["ABONOS"].sum()),
        "NETO": float(df_grouped["NETO"].sum()),
    }
    filas_con_espacios.append(fila_total)

    df_conceptos = pd.DataFrame(filas_con_espacios)

    # PASO 3: REORGANIZACIÓN VISUAL (Negativos arriba)
    df_datos = df.copy()
    df_datos["VALOR"] = valores_numericos
    if tiene_columna_saldo:
        df_datos["SALDO"] = saldos_numericos

    is_negative = df_datos["VALOR"] < 0
    df_datos["_orden_temp"] = 0
    df_datos.loc[~is_negative, "_orden_temp"] = 1

    df_datos_ordenado = df_datos.sort_values(
        by="_orden_temp", kind="stable"
    ).drop(columns=["_orden_temp"])

    # PASO 4: EXPORTACIÓN Y FORMATO EN EXCEL
    with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
        df_datos_ordenado.to_excel(writer, sheet_name="Datos", index=False)
        df_conciliacion.to_excel(writer, sheet_name="Resumen", index=False)
        df_conceptos.to_excel(writer, sheet_name="Conceptos", index=False)

        workbook = writer.book
        FORMATO_MONEDA = '"$"#,##0.00;[Red]("$"#,##0.00);"-"'

        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]

            for col in worksheet.iter_cols(1, worksheet.max_column):
                header_val = (
                    str(col[0].value).upper() if col[0].value else ""
                )

                if header_val in [
                    "VALOR",
                    "SALDO",
                    "CARGOS",
                    "ABONOS",
                    "NETO",
                    "MONTO",
                ]:
                    for cell in col[1:]:
                        if cell.value is not None and isinstance(
                            cell.value, (int, float)
                        ):
                            cell.number_format = FORMATO_MONEDA

    return excel_path


# ==============================================================================
# ORQUESTADOR PRINCIPAL
# ==============================================================================


def ejecutar_proceso_exportacion(pdf_path, output_excel_path=None):
    """Procesa el extracto de Banco de Bogotá directamente sin exigir palabras fijas en el nombre."""
    columnas = [
        "FECHA",
        "COD TRANS",
        "DESCRIPCIÓN",
        "CIUDAD",
        "OFICINA/CANAL",
        "DOCUMENTO",
        "VALOR",
        "SALDO",
    ]

    lineas_plana = extraer_lineas_extractos_pyme(pdf_path)

    if not lineas_plana:
        raise ValueError("No se pudieron extraer movimientos del PDF.")

    if not output_excel_path:
        base_path, _ = os.path.splitext(pdf_path)
        output_excel_path = f"{base_path}_convertido.xlsx"

    base_txt_path, _ = os.path.splitext(output_excel_path)
    txt_path = f"{base_txt_path}_plano.txt"

    # 1. Guardar archivo Plano (.txt)
    encabezado_txt = DELIMITADOR.join(columnas)
    contenido_txt = [encabezado_txt] + lineas_plana

    with open(txt_path, "w", encoding="utf-8") as f:
        f.write("\n".join(contenido_txt))

    # 2. Convertir a Excel
    texto_delimitado_a_excel(lineas_plana, columnas, output_excel_path)

    # 3. Reorganizar y crear pestañas suplementarias
    reorganizar_excel(output_excel_path)

    return output_excel_path